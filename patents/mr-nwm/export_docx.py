#!/usr/bin/env python3
"""Export one patent source using the future-code native Word template."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.table import Table
from docx.text.paragraph import Paragraph
from disclosure_figures import CAPTIONS, diagram

HERE = Path(__file__).resolve().parent
STEMS = ("INVENTION_DISCLOSURE_CN", "PRIOR_ART_SEARCH_CN")
# The native OMML renderer below supports these exact source expressions.
# Refuse silently stale equations if the Markdown source is edited.
EQUATIONS = (
    r"$$u_t=\gamma\odot\operatorname{SiLU}(W_x[s_t;a_t;c]+b_x)+\beta$$",
    r"$$y_t=G_\theta(u_t,m_t)+\operatorname{Norm}(u_t)AB^\mathsf{T}/\sqrt{r}$$",
    r"$$s_{t+1}=s_t+W_\delta y_t+b_\delta$$",
)
TITLE = "一种无线接入网多时间尺度数据协同处理与资源控制方法及装置"
LEGACY_TITLE = "一种多模态多速率无线网络世界模型控制方法"
FORMAT_PARTS = (
    "word/styles.xml",
    "word/theme/theme1.xml",
    "word/numbering.xml",
    "word/settings.xml",
    "word/fontTable.xml",
    "word/webSettings.xml",
)
FORMAT_HASHES = {
    "disclosure.docx": "eb92d21e7858dc1e94323e44224e607d48cb5b82",
    "search.docx": "295ca826e2467a181b12b735b521ebbbd77c4a9c",
}


def font_file(value=None):
    options = (
        [Path(value)]
        if value
        else [
            Path("/usr/local/share/fonts/noto-cjk/NotoSansCJKsc-Regular.otf"),
            Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        ]
    )
    for p in options:
        if p.is_file():
            return p
    raise SystemExit("Pass --font-file with a Chinese font for the diagrams.")


def add_link(p, url):
    url = url.rstrip("。，；、.,;")
    if "3gpp.org/news-events/3gpp-news/ai-ml-2025" in url:
        label = "3GPP AI/ML技术介绍"
    elif "3gpp.org/3gpp-groups/radio-access-networks-ran" in url:
        label = "3GPP RAN职责"
    elif "/patent/" in url:
        label = re.search(r"/patent/([^/]+)", url).group(1) + " 文献页"
    elif "/abs/" in url:
        label = "arXiv " + url.rsplit("/", 1)[1]
    elif "cnipa.gov.cn" in url:
        label = "CNIPA 专利法" if "/2020/" in url else "CNIPA 局令第84号"
    elif "/commit/" in url:
        label = "项目版本 " + url.rsplit("/", 1)[1][:12]
    else:
        label = url.split("//", 1)[-1]
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    pr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    pr.append(style)
    run.append(pr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    p._p.append(link)


def inline(p, value, properties=None):
    position = 0
    for m in re.finditer(r"https?://[^\s。；）)]+|\*\*[^*]+\*\*|`[^`]+`", value):
        if m.start() > position:
            p.add_run(value[position : m.start()])
        s = m.group()
        if s.startswith("http"):
            add_link(p, s)
        else:
            run = p.add_run(s[2:-2] if s.startswith("**") else s[1:-1])
            if s.startswith("**"):
                run.bold = True
        position = m.end()
    p.add_run(value[position:])
    if properties is not None:
        for run in p.runs:
            run._r.insert(0, deepcopy(properties))


def reset(element):
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)


def math_run(text):
    run = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    run.append(t)
    return run


def script(base, suffix, superscript=False):
    el = OxmlElement("m:sSup" if superscript else "m:sSub")
    e = OxmlElement("m:e")
    e.append(math_run(base))
    el.append(e)
    e = OxmlElement("m:sup" if superscript else "m:sub")
    e.append(math_run(suffix))
    el.append(e)
    return el


def equation(paragraph, index):
    """Native, editable OMML with actual subscript/superscript/fraction/radical nodes."""
    container = OxmlElement("m:oMathPara")
    expr = OxmlElement("m:oMath")
    container.append(expr)
    r = math_run
    sub = script
    if index == 0:
        nodes = [
            sub("u", "t"),
            r(" = γ ⊙ SiLU("),
            sub("W", "x"),
            r("["),
            sub("s", "t"),
            r("; "),
            sub("a", "t"),
            r("; c] + "),
            sub("b", "x"),
            r(") + β"),
        ]
    elif index == 1:
        frac = OxmlElement("m:f")
        numerator = OxmlElement("m:num")
        for n in [r("Norm("), sub("u", "t"), r(") A "), script("B", "T", True)]:
            numerator.append(n)
        denominator = OxmlElement("m:den")
        radical = OxmlElement("m:rad")
        props = OxmlElement("m:radPr")
        hide = OxmlElement("m:degHide")
        hide.set(qn("m:val"), "1")
        props.append(hide)
        radical.append(props)
        radical.append(OxmlElement("m:deg"))
        val = OxmlElement("m:e")
        val.append(r("r"))
        radical.append(val)
        denominator.append(radical)
        frac.append(numerator)
        frac.append(denominator)
        nodes = [
            sub("y", "t"),
            r(" = "),
            sub("G", "θ"),
            r("("),
            sub("u", "t"),
            r(", "),
            sub("m", "t"),
            r(") + "),
            frac,
        ]
    else:
        nodes = [
            sub("s", "t+1"),
            r(" = "),
            sub("s", "t"),
            r(" + "),
            sub("W", "δ"),
            sub("y", "t"),
            r(" + "),
            sub("b", "δ"),
        ]
    for n in nodes:
        expr.append(n)
    paragraph._p.append(container)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1


class Native:
    def __init__(self, path):
        self.path = path
        data = path.read_bytes()
        blob = hashlib.sha1(f"blob {len(data)}\0".encode() + data).hexdigest()
        assert blob == FORMAT_HASHES[path.name], "Native template bytes changed"
        self.doc = Document(path)
        self.prototypes = {}
        self.heading_sequence = []
        for p in self.doc.paragraphs:
            self.prototypes.setdefault(p.style.name, deepcopy(p._p))
            if re.fullmatch(r"Heading [123]", p.style.name):
                self.heading_sequence.append(p.style.name)
        self.tables = [deepcopy(t._tbl) for t in self.doc.tables]
        for e in list(self.doc._element.body):
            if e.tag != qn("w:sectPr"):
                self.doc._element.body.remove(e)

    def append(self, element):
        self.doc._element.body.insert(len(self.doc._element.body) - 1, element)

    def paragraph(self, style, value=""):
        element = deepcopy(self.prototypes[style])
        reset(element)
        self.append(element)
        p = Paragraph(element, self.doc._body)
        inline(p, value)
        if self.path.name == "disclosure.docx":
            size = {"Title": 24, "Heading 1": 16, "Heading 2": 13,
                    "Heading 3": 11.5, "Image Caption": 10}.get(style, 11)
            p.paragraph_format.line_spacing = Pt(size + 4.5)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(5)
            if style.startswith("Heading"):
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.keep_with_next = True
            for run in p.runs:
                run.font.name = "Noto Sans CJK SC"
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
                run.font.size = Pt(size)
                run.font.color.rgb = RGBColor(0, 0, 0)
                if style.startswith("Heading"):
                    run.bold = True
        return p

    def table(self, index, rows):
        if self.path.name != "disclosure.docx":
            return self.legacy_table(index, rows)
        widths = [2100, 7260] if len(rows[0]) == 2 else [1800, 4800, 2760]
        assert len(widths) == len(rows[0]), "Unsupported disclosure table schema"
        table = self.doc.add_table(rows=0, cols=len(widths))
        table.autofit = False
        for col, width in zip(table.columns, widths):
            col.width = Inches(width / 1440)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:" + side)
            for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                el.set(qn("w:" + key), value)
            borders.append(el)
        table._tbl.tblPr.append(borders)
        for i, record in enumerate(rows):
            row = table.add_row()
            props = row._tr.get_or_add_trPr()
            props.append(OxmlElement("w:cantSplit"))
            if i == 0:
                props.append(OxmlElement("w:tblHeader"))
            for cell, value, width in zip(row.cells, record, widths):
                cell.width = Inches(width / 1440)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                margin = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    el = OxmlElement("w:" + side)
                    el.set(qn("w:w"), "40" if side in ("top", "bottom") else "90")
                    el.set(qn("w:type"), "dxa")
                    margin.append(el)
                tcpr.append(margin)
                if i == 0:
                    fill = OxmlElement("w:shd")
                    fill.set(qn("w:fill"), "E6E6E6")
                    tcpr.append(fill)
                p = cell.paragraphs[0]
                inline(p, value)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(13)
                p.paragraph_format.space_before = Pt(0)
                for run in p.runs:
                    run.font.name = "Noto Sans CJK SC"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.bold = i == 0

    def legacy_table(self, index, rows):
        element = deepcopy(self.tables[index])
        table = Table(element, self.doc._body)
        assert len(table.columns) == len(rows[0])
        widths = [int(col.get(qn("w:w"))) for col in element.tblGrid]
        # The source grid is authoritative. Explicit widths prevent the
        # empty-cell prototypes from collapsing header cells in LibreOffice.
        layout = element.tblPr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            element.tblPr.append(layout)
        layout.set(qn("w:type"), "fixed")
        total = element.tblPr.find(qn("w:tblW"))
        if total is None:
            total = OxmlElement("w:tblW")
            element.tblPr.append(total)
        total.set(qn("w:type"), "dxa")
        total.set(qn("w:w"), str(sum(widths)))
        head = deepcopy(table.rows[0]._tr)
        body = deepcopy(table.rows[1]._tr)
        for row in list(element.findall(qn("w:tr"))):
            element.remove(row)
        for i, record in enumerate(rows):
            tr = deepcopy(head if i == 0 else body)
            tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell, value, width in zip(tr.findall(qn("w:tc")), record, widths):
                cell_props = cell.find(qn("w:tcPr"))
                if cell_props is None:
                    cell_props = OxmlElement("w:tcPr")
                    cell.insert(0, cell_props)
                cell_width = cell_props.find(qn("w:tcW"))
                if cell_width is None:
                    cell_width = OxmlElement("w:tcW")
                    cell_props.append(cell_width)
                cell_width.set(qn("w:type"), "dxa")
                cell_width.set(qn("w:w"), str(width))
                proto = deepcopy(cell.find(qn("w:p")))
                r = proto.find(qn("w:r"))
                props = deepcopy(r.find(qn("w:rPr"))) if r is not None else None
                for child in list(cell):
                    if child.tag != qn("w:tcPr"):
                        cell.remove(child)
                reset(proto)
                cell.append(proto)
                inline(Paragraph(proto, self.doc._body), value, props)
            element.append(tr)
        self.append(element)

    def title(self, search, headings, toc):
        p = self.paragraph("Title")
        p.add_run("现有技术检索报告" if search else "发明专利技术交底书").bold = True
        p.add_run().add_break()
        if search:
            p.add_run("Prior Art Search Report").add_break()
            p.add_run("多模态多速率无线网络世界模型控制")
        else:
            p.add_run(TITLE)
        if not search:
            for run in p.runs:
                run.font.name = "Noto Sans CJK SC"
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
                run.font.size = Pt(22)
                run.font.color.rgb = RGBColor(0, 0, 0)
        self.paragraph("Author", "中国移动通信有限公司研究院")
        self.paragraph("Date", "2026-10-16" if search else "2026-10-22")
        p = self.doc.add_paragraph("目录")
        p.paragraph_format.space_before = Pt(15)
        p.paragraph_format.space_after = Pt(6)
        p.runs[0].bold = True
        for h in headings:
            label = h.split(" / ")[0]
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.tab_stops.add_tab_stop(
                Inches(5.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
            )
            run = p.add_run(label + ("\t" + str(toc[h]) if h in toc else ""))
            run.font.size = Pt(11)
        self.doc.add_page_break()


def preserve_parts(target, template):
    with zipfile.ZipFile(target) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    with zipfile.ZipFile(template) as z:
        for n in FORMAT_PARTS:
            entries[n] = z.read(n)
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries.items():
            z.writestr(name, data)


def export(stem, font, toc):
    search = stem == STEMS[1]
    native = Native(HERE / "formats" / ("search.docx" if search else "disclosure.docx"))
    lines = (HERE / (stem + ".md")).read_text().splitlines()
    roots = [line[2:] for line in lines if line.startswith("# ")]
    native.title(search, roots, toc.get(stem, {}))
    i = next(i for i, line in enumerate(lines) if line.startswith("# "))
    headings = []
    table_count = figure_count = equation_count = 0
    first = True
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line.startswith("```mermaid"):
            while i < len(lines) and not lines[i].startswith("```"):
                i += 1
            i += 1
            figure_count += 1
            p = native.paragraph("Captioned Figure")
            p.paragraph_format.line_spacing = 1
            pic = p.add_run().add_picture(str(diagram(figure_count, font)), width=Inches(6.3))
            pic._inline.docPr.set("descr", CAPTIONS[figure_count - 1])
            p.paragraph_format.keep_with_next = True
            caption = native.paragraph("Image Caption", CAPTIONS[figure_count - 1])
            caption.paragraph_format.keep_with_next = False
            continue
        if line.startswith("$$"):
            assert line == EQUATIONS[equation_count], "Update native equation renderer for changed source"
            p = native.paragraph("First Paragraph")
            equation(p, equation_count)
            equation_count += 1
            continue
        if line.startswith("|"):
            block = [line]
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = [
                [v.strip() for v in row.strip("|").split("|")]
                for row in block
                if not re.fullmatch(r"\|[\s:|\-]+\|", row)
            ]
            assert all(len(row) == len(rows[0]) for row in rows)
            native.table(table_count, rows)
            table_count += 1
            first = False
            continue
        match = re.fullmatch(r"(#{1,3})\s+(.+)", line)
        if match:
            style = f"Heading {len(match.group(1))}"
            paragraph = native.paragraph(style, match.group(2))
            if not search and match.group(2) in ("实施例 3 信息交换与6G标准化适配", "权利要求书"):
                paragraph.paragraph_format.page_break_before = True
            headings.append(style)
            first = True
            continue
        paragraph = native.paragraph("First Paragraph" if first else "Body Text", line)
        if not search and re.match(r"^\d+\.", line):
            paragraph.paragraph_format.keep_together = True
        first = False
    if search:
        assert headings == native.heading_sequence, "Heading levels/order differ from native template"
    else:
        expected_roots = ["发明名称", "申请人信息", "发明人信息", "技术领域", "背景技术", "发明内容", "附图说明", "具体实施方式", "技术效果", "权利要求书", "说明书摘要"]
        assert roots == expected_roots, "Root sections/order differ from native template"
    assert table_count == len(native.tables)
    assert figure_count == (0 if search else 5)
    assert equation_count == (0 if search else 3)
    props = native.doc.core_properties
    props.author = props.last_modified_by = ""
    props.title = "现有技术检索报告" if search else "发明专利技术交底书"
    props.subject = LEGACY_TITLE if search else TITLE
    props.comments = ""
    props.created = props.modified = datetime(2026, 10, 16 if search else 22, tzinfo=timezone.utc)
    section = native.doc.sections[0]
    if not search:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(0.72)
        section.left_margin = section.right_margin = Inches(1)
        section.footer_distance = Inches(0.3)
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for val in ["第 ", "PAGE", " 页"]:
        if val == "PAGE":
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE")
            p._p.append(field)
        else:
            p.add_run(val).font.size = Pt(10)
    target = HERE / (stem + ".docx")
    native.doc.save(target)
    preserve_parts(target, native.path)
    with zipfile.ZipFile(target) as z, zipfile.ZipFile(native.path) as t:
        assert all(z.read(part) == t.read(part) for part in FORMAT_PARTS)
    return target


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--font-file")
    parser.add_argument("--only", choices=("disclosure", "search"), default="disclosure")
    parser.add_argument("--toc-map", type=Path)
    args = parser.parse_args()
    toc = json.loads(args.toc_map.read_text()) if args.toc_map else {}
    font = font_file(args.font_file)
    stem = STEMS[0 if args.only == "disclosure" else 1]
    print(export(stem, font, toc))


if __name__ == "__main__":
    main()
